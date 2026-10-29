#!/usr/bin/env python3
# =============================================================
# Nombre del archivo: build_byid_robot_match.py
# Descripción: Arma un Excel para relacionar las vacas de ByID (IDs
#   propios, agrupados a mano) con el número de animal del robot.
#   - El ID del robot de cada foto sale de dataset_del.csv (registro
#     del robot más cercano en tiempo a la foto: vacaId, foundDate,
#     difImagen). Es una sugerencia, no la verdad.
#   - El BCS sale de <labels_dir>/<BCS>/img.jpg (etiquetas nuevas).
#   - Generalidades de Ordeño (reporte diario) se usa para validar:
#     ¿el animal estaba en ordeño ese día y tuvo al menos tantos
#     ordeños como visitas fotografiadas?
#   Hojas: Instrucciones, Vacas_ByID (una fila por vaca, con columna
#   para confirmar el ID del robot), Fotos (una fila por foto),
#   IDs_robot (en qué grupos de ByID aparece cada ID del robot).
#
# Uso (requiere openpyxl):
#   python3 build_byid_robot_match.py --ids_dir Dataset/ByID --labels_dir Dataset/Images \
#       --csv dataset_del.csv --ordeno_dir GeneralidadesOrdeno --out Dataset/byid_vs_robot.xlsx
#   Con --out <carpeta> (sin .xlsx) escribe vacas_byid.csv, fotos.csv e ids_robot.csv,
#   sin las columnas de fórmula.
# =============================================================

import argparse, csv, re
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VISIT_GAP_MIN = 10  # fotos separadas por más de esto = visitas distintas
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
RED_FILL = PatternFill("solid", fgColor="F8CBAD")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")


def parse_timedelta_min(s):
    """'0 days 00:14:31' / '-1 days +23:50:00' (formato pandas) -> minutos."""
    m = re.match(r"(-?\d+) days \+?(\d+):(\d+):(\d+)", s or "")
    if not m:
        return None
    d, h, mi, se = map(int, m.groups())
    return round(d * 1440 + h * 60 + mi + se / 60, 1)


def load_ordeno(ordeno_dir):
    """{fecha: {animal: fila}} y {fecha: {animal: ordeños de ese día}}."""
    info, milkings = {}, {}
    for f in Path(ordeno_dir).glob("*.csv"):
        d = datetime.strptime(re.search(r"_(\d{8})", f.name).group(1), "%Y%m%d").date()
        rows = [r for r in csv.DictReader(open(f, encoding="latin1"), delimiter=";") if r["Número del animal"]]
        info[d] = {r["Número del animal"]: r for r in rows}
        # El archivo del día d reporta los ordeños de "ayer" (d-1)
        milkings[d - timedelta(days=1)] = {r["Número del animal"]: int(r["Ordeños Ayer"])
                                           for r in rows if r["Ordeños Ayer"]}
    return info, milkings


def load_photos(ids_dir, labels_dir, csv_path):
    labels = {}
    for p in Path(labels_dir).rglob("*.jpg"):
        try:
            labels.setdefault(p.name, float(p.parent.name))
        except ValueError:
            pass
    robot = {r["ImgName"]: r for r in csv.DictReader(open(csv_path, newline=""))}
    photos, seen = [], set()
    for p in sorted(Path(ids_dir).rglob("*.jpg")):
        name = p.name.replace(" - Copy", "")
        rel = p.relative_to(ids_dir).parts
        group = int(rel[0]) if rel[0].isdigit() else "/".join(rel[:-1])
        if (group, name) in seen:
            continue
        seen.add((group, name))
        r = robot.get(name[:-4], {})
        photos.append({
            "byid": group, "img": name,
            "ts": datetime.strptime(name[:19], "%Y-%m-%d-%H-%M-%S"),
            "cam": name[20:24],
            "bcs": labels.get(name), "bcs_csv": float(r["BCS"]) if r else None,
            "robot": int(r["vacaId"]) if r else None,
            "robot_ts": datetime.fromisoformat(r["foundDate"]) if r else None,
            "dif_min": parse_timedelta_min(r.get("difImagen")) if r else None,
            "del": float(r["DEL"]) if r and r["DEL"] else None,
            "flagMatch": r.get("flagMatch", ""), "valido": r.get("valido", ""),
        })
    return photos


def visits_per_day(photos):
    out, last = Counter(), None
    for t in sorted(p["ts"] for p in photos):
        if last is None or (t - last).total_seconds() > VISIT_GAP_MIN * 60:
            out[t.date()] += 1
        last = t
    return out


def incompatible_days(robot_id, visits, info, milkings):
    """Días con fotos en que el animal no estaba en ordeño o tuvo menos ordeños que visitas."""
    bad = []
    for d, n in visits.items():
        if d not in info:
            continue
        if str(robot_id) not in info[d]:
            bad.append(f"{d:%d-%b} ausente")
        elif d in milkings and milkings[d].get(str(robot_id), 0) < n:
            bad.append(f"{d:%d-%b} {n} visitas > {milkings[d].get(str(robot_id), 0)} ordeños")
    return bad


def style_sheet(ws, widths, input_cols=()):
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
    for col in input_cols:
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col).fill = INPUT_FILL
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ids_dir", required=True)
    ap.add_argument("--labels_dir", required=True)
    ap.add_argument("--csv", required=True)
    ap.add_argument("--ordeno_dir", required=True)
    ap.add_argument("--out", required=True, help="Archivo .xlsx, o una carpeta para escribir los CSV.")
    args = ap.parse_args()
    as_csv = not args.out.endswith(".xlsx")
    tables = {}  # hoja -> (encabezado, filas sin fórmulas) para la salida CSV

    info, milkings = load_ordeno(args.ordeno_dir)
    photos = load_photos(Path(args.ids_dir), args.labels_dir, args.csv)
    by_group = defaultdict(list)
    for p in photos:
        by_group[p["byid"]].append(p)
    groups = sorted(by_group, key=lambda g: (isinstance(g, str), g if isinstance(g, int) else 0, str(g)))

    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"

    # ---------- Vacas_ByID ----------
    wv = wb.create_sheet("Vacas_ByID")
    head = ["ID ByID", "Fotos", "Fotos con ID robot", "Días con fotos", "Primera foto", "Última foto",
            "BCS mediana", "BCS mín", "BCS máx",
            "ID robot sugerido", "% fotos que lo apoyan", "2º candidato", "% 2º",
            "Todos los candidatos (ID: fotos)", "|dif| media al sugerido (min)",
            "Días incompatibles del sugerido (Generalidades)",
            "ID robot confirmado", "¿ID repetido?", "Notas"]
    wv.append(head)
    for g in groups:
        ps = by_group[g]
        cands = Counter(p["robot"] for p in ps if p["robot"] is not None)
        top = cands.most_common(2)
        n_rob = sum(cands.values())
        bcs = sorted(p["bcs"] for p in ps if p["bcs"] is not None)
        sug = top[0][0] if top else None
        difs = [abs(p["dif_min"]) for p in ps if p["robot"] == sug and p["dif_min"] is not None]
        bad = incompatible_days(sug, visits_per_day(ps), info, milkings) if sug else []
        r = wv.max_row + 1
        wv.append([
            g, len(ps), n_rob, len({p["ts"].date() for p in ps}),
            min(p["ts"] for p in ps), max(p["ts"] for p in ps),
            bcs[len(bcs) // 2] if bcs else None, bcs[0] if bcs else None, bcs[-1] if bcs else None,
            sug, top[0][1] / n_rob if top else None,
            top[1][0] if len(top) > 1 else None, top[1][1] / n_rob if len(top) > 1 else None,
            ", ".join(f"{k}: {v}" for k, v in cands.most_common()),
            round(sum(difs) / len(difs), 1) if difs else None,
            "; ".join(bad) if bad else ("" if sug else "sin ID robot"),
            None,
            f'=IF(Q{r}="","",IF(COUNTIF($Q$2:$Q$200,Q{r})>1,"REPETIDO",""))',
            None,
        ])
        tables.setdefault("vacas_byid", (head[:17] + head[18:], []))[1].append(
            [c.value for c in wv[r]][:17] + [None])
    nv = wv.max_row
    style_sheet(wv, [14, 7, 9, 8, 17, 17, 8, 7, 7, 11, 11, 11, 8, 40, 11, 42, 13, 11, 30], input_cols=(17, 19))
    for row in range(2, nv + 1):
        for c in (5, 6):
            wv.cell(row=row, column=c).number_format = "dd-mmm-yy hh:mm"
        for c in (11, 13):
            wv.cell(row=row, column=c).number_format = "0%"
    wv.conditional_formatting.add(f"K2:K{nv}", CellIsRule(operator="lessThan", formula=["0.6"], fill=RED_FILL))
    wv.conditional_formatting.add(f"K2:K{nv}", CellIsRule(operator="greaterThanOrEqual", formula=["0.8"], fill=GREEN_FILL))
    wv.conditional_formatting.add(f"P2:P{nv}", FormulaRule(formula=[f'LEN(P2)>0'], fill=RED_FILL))
    wv.conditional_formatting.add(f"R2:R{nv}", FormulaRule(formula=[f'R2="REPETIDO"'], fill=RED_FILL))
    wv["Q1"].comment = Comment("Escribe aquí el número de animal del robot que confirmes para esta vaca "
                               "(p.ej. 6003). Déjalo vacío si no estás seguro.", "Claude")
    wv["K1"].comment = Comment("Fracción de las fotos de esta vaca (con ID robot) cuyo registro más "
                               "cercano en dataset_del.csv es el ID sugerido. Verde >= 80%, rojo < 60%.", "Claude")
    wv["P1"].comment = Comment(f"Revisado contra Generalidades de Ordeño: días en que el ID sugerido no "
                               f"estaba en ordeño, o tuvo menos ordeños que visitas fotografiadas "
                               f"(fotos separadas > {VISIT_GAP_MIN} min = visitas distintas).", "Claude")

    # ---------- Fotos ----------
    wf = wb.create_sheet("Fotos")
    fotos_head = ["ID ByID", "Imagen", "Fecha/hora foto", "Cámara", "BCS (nuevo)", "BCS (CSV viejo)",
               "ID robot (CSV)", "Hora registro robot", "Dif. foto − registro (min)",
               "DEL (CSV)", "flagMatch", "valido",
               "En ordeño ese día", "Ordeños ese día", "Visitas fotografiadas ese día",
               "ID robot confirmado (de Vacas_ByID)", "¿Coincide?"]
    wf.append(fotos_head)
    visits = {g: visits_per_day(by_group[g]) for g in groups}
    for g in groups:
        for p in sorted(by_group[g], key=lambda p: p["ts"]):
            d = p["ts"].date()
            rid = str(p["robot"]) if p["robot"] is not None else None
            present = ("sí" if rid in info[d] else "no") if rid and d in info else None
            r = wf.max_row + 1
            wf.append([
                g, p["img"], p["ts"], p["cam"], p["bcs"], p["bcs_csv"],
                p["robot"], p["robot_ts"], p["dif_min"], p["del"], p["flagMatch"], p["valido"],
                present, milkings.get(d, {}).get(rid) if rid else None, visits[g][d],
                f'=IFERROR(IF(INDEX(Vacas_ByID!$Q$2:$Q${nv},MATCH(A{r},Vacas_ByID!$A$2:$A${nv},0))="","",'
                f'INDEX(Vacas_ByID!$Q$2:$Q${nv},MATCH(A{r},Vacas_ByID!$A$2:$A${nv},0))),"")',
                f'=IF(OR(P{r}="",G{r}=""),"",IF(P{r}=G{r},"sí","no"))',
            ])
            tables.setdefault("fotos", (fotos_head[:15], []))[1].append([c.value for c in wf[r]][:15])
    nf = wf.max_row
    style_sheet(wf, [14, 34, 17, 7, 8, 8, 10, 17, 10, 7, 9, 8, 9, 9, 10, 13, 9])
    for row in range(2, nf + 1):
        wf.cell(row=row, column=3).number_format = "dd-mmm-yy hh:mm:ss"
        wf.cell(row=row, column=8).number_format = "dd-mmm-yy hh:mm"
        wf.cell(row=row, column=16).font = Font(name=FONT, color="008000")
    wf.conditional_formatting.add(f"Q2:Q{nf}", FormulaRule(formula=['Q2="no"'], fill=RED_FILL))
    wf.conditional_formatting.add(f"M2:M{nf}", FormulaRule(formula=['M2="no"'], fill=RED_FILL))
    wf.conditional_formatting.add(f"I2:I{nf}", FormulaRule(formula=['ABS(I2)>30'], fill=RED_FILL))

    # ---------- IDs_robot ----------
    wr = wb.create_sheet("IDs_robot")
    mid = sorted(info)[len(info) // 2]
    rob_head = ["ID robot", "Fotos asignadas (CSV)", "Grupos ByID donde aparece (ByID: fotos)", "Nº grupos ByID",
               f"Lactación ({mid:%d-%b})", f"Grupo ({mid:%d-%b})", f"DEL ({mid:%d-%b})",
               f"Estado repro ({mid:%d-%b})", "Confirmado para ByID"]
    wr.append(rob_head)
    rob_groups = defaultdict(Counter)
    for p in photos:
        if p["robot"] is not None:
            rob_groups[p["robot"]][p["byid"]] += 1
    for rid in sorted(rob_groups):
        gi = info[mid].get(str(rid), {})
        r = wr.max_row + 1
        wr.append([rid, sum(rob_groups[rid].values()),
                   ", ".join(f"{g}: {n}" for g, n in rob_groups[rid].most_common()), len(rob_groups[rid]),
                   gi.get("Nº Lactación"), gi.get("Número del grupo"),
                   int(gi["Días en ordeño"]) if gi.get("Días en ordeño") else None,
                   gi.get("Estado de la reproducción"),
                   f'=IFERROR(INDEX(Vacas_ByID!$A$2:$A${nv},MATCH(A{r},Vacas_ByID!$Q$2:$Q${nv},0)),"")'])
        tables.setdefault("ids_robot", (rob_head[:8], []))[1].append([c.value for c in wr[r]][:8])
    style_sheet(wr, [10, 11, 60, 9, 11, 10, 9, 13, 12])
    for row in range(2, wr.max_row + 1):
        wr.cell(row=row, column=9).font = Font(name=FONT, color="008000")

    # ---------- Instrucciones ----------
    n_rob = sum(p["robot"] is not None for p in photos)
    lines = [
        ("Relación ByID ↔ ID del robot", True),
        ("", False),
        ("Qué es cada hoja", True),
        ("Vacas_ByID: una fila por vaca de ByID (tus IDs propios). Aquí decides el ID del robot.", False),
        ("Fotos: una fila por foto, con el ID del robot que le asignó dataset_del.csv por cercanía en tiempo.", False),
        ("IDs_robot: cada ID del robot y en cuántos grupos de ByID aparece (lo ideal es 1).", False),
        ("", False),
        ("Qué llenar (celdas amarillas)", True),
        ("Vacas_ByID, columna Q 'ID robot confirmado': el número de animal que confirmes (p.ej. 6003). Vacío = sin decidir.", False),
        ("Vacas_ByID, columna S 'Notas': libre (p.ej. 'mezcla de dos vacas', 'revisar arete').", False),
        ("Al llenar Q se actualizan solos: '¿ID repetido?' (Vacas_ByID), '¿Coincide?' (Fotos) y 'Confirmado para ByID' (IDs_robot).", False),
        ("", False),
        ("Cómo leer la sugerencia", True),
        ("ID robot sugerido = el ID más frecuente entre las fotos de la vaca. '% fotos que lo apoyan': verde >= 80%, rojo < 60%.", False),
        ("Días incompatibles: días en que el ID sugerido no estaba en ordeño, o tuvo menos ordeños que visitas fotografiadas"
         f" (fotos a más de {VISIT_GAP_MIN} min = visitas distintas). Si hay alguno, la sugerencia es dudosa o el grupo mezcla vacas.", False),
        ("Dif. foto − registro: minutos entre la foto y el registro del robot usado. En rojo si es mayor a 30 min.", False),
        ("", False),
        ("Limitaciones", True),
        (f"dataset_del.csv solo trae {len(rob_groups)} IDs del robot distintos para {len(groups)} grupos de ByID: "
         "muchas vacas reales NO tienen su ID entre los candidatos. Un ID sugerido no es un ID confirmado.", False),
        (f"Fotos: {len(photos)} en total, {n_rob} con ID robot en dataset_del.csv.", False),
        ("Fuentes: Dataset/ByID (grupos), Dataset/Images (BCS nuevo), dataset_del.csv (ID robot y hora), "
         "GeneralidadesOrdeno (presencia y ordeños por día).", False),
        (f"Generado el {datetime.now():%Y-%m-%d} con Modeling/Support Code/build_byid_robot_match.py.", False),
    ]
    for text, bold in lines:
        ins.append([text])
        ins.cell(row=ins.max_row, column=1).font = Font(name=FONT, bold=bold, size=14 if text.startswith("Relación") else 11)
    ins.column_dimensions["A"].width = 140

    if as_csv:
        write_csvs(tables, Path(args.out))
    else:
        wb.calculation.fullCalcOnLoad = True
        wb.active = 1
        Path(args.out).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.out)

    strong = sum(1 for row in wv.iter_rows(min_row=2, values_only=True) if row[10] is not None and row[10] >= 0.8)
    clean = sum(1 for row in wv.iter_rows(min_row=2, values_only=True) if row[9] is not None and not row[15])
    print(f"{len(groups)} grupos, {len(photos)} fotos ({n_rob} con ID robot), {len(rob_groups)} IDs robot")
    print(f"Sugerencia con >=80% de apoyo: {strong} | sin días incompatibles: {clean}")
    print("Salida:", args.out)


def fmt_cell(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, float):
        return round(v, 3)
    return "" if v is None else v


def write_csvs(tables, out_dir):
    """Un CSV por hoja (UTF-8 con BOM para que Excel respete los acentos), sin columnas de fórmula."""
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, (head, rows) in tables.items():
        with open(out_dir / f"{name}.csv", "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(head)
            w.writerows([fmt_cell(v) for v in row] for row in rows)


if __name__ == "__main__":
    main()
